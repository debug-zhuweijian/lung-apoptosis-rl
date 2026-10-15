"""Small constructed workbooks test parsing, not biological validity."""
from openpyxl import Workbook
import pandas as pd
import pytest

from scripts.recover_apotracker_reference import read_trajectories, treatment_events


def make_book(path, bad=None):
    w = Workbook()
    w.remove(w.active)
    specs = [("nuclei_count", "Nuclei counts", [100, 90]),
             ("fraction_Green", "ApoTracker Green only", [20, 30]),
             ("fraction_GreenRed", "ApoTracker Green & propidium iodide", [10, 20]),
             ("fraction_Red", "Propidium iodide only", [5, 10])]
    for suffix, title, values in specs:
        s = w.create_sheet("Apoptosis_" + suffix)
        s.append([title]); s.append([]); s.append([])
        s.append(["Cell line", "Condition", "Treatment", "Replicate", 3, 6.4])
        if bad == "overflow" and suffix == "fraction_Green":
            values = [99, 99]
        if bad == "fractional_count" and suffix == "nuclei_count":
            values = [100.5, 90]
        if bad == "clock" and suffix == "fraction_Red":
            s.cell(4, 6, 7)
        s.append(["A549", "S-", "Trimetazidine_400uM", 1, *values])
        if bad == "duplicate" and suffix == "fraction_Red":
            s.append(["A549", "S-", "Trimetazidine_400uM", 1, *values])
    w.save(path)


def test_disjoint_source_channels_not_double_subtracted(tmp_path):
    p = tmp_path / "input.xlsx"; make_book(p)
    result = read_trajectories(p)
    assert result.time_h.tolist() == [3, 6.4]
    assert result.ps_positive_percent.tolist() == [30, 50]
    assert result.pi_positive_percent.tolist() == [15, 30]
    assert result.double_negative_percent.tolist() == [65, 40]
    assert not result.has_matched_RNA.any()
    assert result.assay.eq("ApoTracker_Green_PI_imaging").all()


@pytest.mark.parametrize("bad,match", [("overflow", "sum exceeds"),
    ("fractional_count", "integer"), ("clock", "keys differ"), ("duplicate", "Duplicate")])
def test_invalid_source_never_silently_clipped_or_joined(tmp_path, bad, match):
    p = tmp_path / "input.xlsx"; make_book(p, bad)
    with pytest.raises(ValueError, match=match):
        read_trajectories(p)


def test_joint_drugs_are_concurrent_events_not_sequence():
    events = treatment_events("TrimetazidineEtomoxir_1000uM10uM")
    assert [e["dose_uM"] for e in events] == [1000, 10]
    assert all(e["start_h"] == 0 for e in events)
    with pytest.raises(ValueError, match="Unrecognized"):
        treatment_events("Trimetazidine_1mM")
