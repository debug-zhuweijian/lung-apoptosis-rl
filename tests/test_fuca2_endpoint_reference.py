import openpyxl
import pytest
from scripts.recover_fuca2_endpoint_reference import extract, DOUBLE, DOSES


def workbook():
    w=openpyxl.Workbook();s=w.active;s.title='figS7';s['C17']='figS7.b'
    for cell,r,name in [('A549',19,'A549'),('NCI-H1299',23,'H1299')]:
        s.cell(r-1,3,name)
        for i,d in enumerate(DOSES):
            s.cell(r-1,4+3*i,f'{d} μm')
            for j in range(3):s.cell(r,4+3*i+j,DOUBLE[cell][i]+j*.1)
    return w


def test_units_backgrounds_and_no_fake_cells():
    rows,reps=extract(workbook())
    assert len(rows)==36 and len(reps)==12
    assert len(rows[rows.cell_background=='A549'])==18
    assert not rows.total_annexin_training_eligible.any()
    assert rows.biological_repeat_id.isna().all()
    assert not rows.observations_are_single_cells.any()
    assert (rows.exposure_hours==48).all()
    assert reps.iloc[0].representative_total_annexin_percent==pytest.approx(9.74)
    assert rows.iloc[0].value_percent==6.78


@pytest.mark.parametrize('address,value',[('C18','H1299'),('D18','0 nM'),('D19',9.74),('E19',101),('C17','figS7.c')])
def test_binding_and_definition_changes_fail(address,value):
    w=workbook();w['figS7'][address]=value
    with pytest.raises(ValueError):extract(w)
