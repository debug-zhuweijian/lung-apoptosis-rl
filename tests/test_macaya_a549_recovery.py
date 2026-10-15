import openpyxl
import pytest
from scripts.recover_macaya_a549_endpoints import extract_a549


def workbook():
    w=openpyxl.Workbook();w.remove(w.active)
    for name,start,second in [('Fig Suppl 2',12,'Lest'),('Fig Suppl 3',27,'Mido')]:
        s=w.create_sheet(name);s.cell(start-2,15,'A549 - Annexin V positive cells (%)')
        for j,label in enumerate(['Control' if start==12 else 'Ctrl','Tram',second,'Combo']):
            s.cell(start-1,15+j,label)
            for i in range(3):s.cell(start+i,15+j,i+j)
    w['Fig Suppl 3'].cell(29,15).value=None;w['Fig Suppl 3'].cell(29,16).value=None
    return w


def test_missing_source_replicates_and_combinations_preserved():
    r=extract_a549(workbook());assert len(r)==24 and sum(x['missing'] for x in r)==2
    assert sum(len(x['components'])==2 for x in r)==6
    assert r[0]['percent']==0 and not r[0]['missing']


@pytest.mark.parametrize('case',['background','order','percent'])
def test_wrong_sheet_binding_is_rejected(case):
    w=workbook();s=w['Fig Suppl 2']
    if case=='background':s.cell(10,15,'H1792')
    elif case=='order':s.cell(11,15,'Tram')
    else:s.cell(12,15,101)
    with pytest.raises(ValueError):extract_a549(w)


def test_changed_protocol_source_is_rejected(tmp_path):
    from scripts.recover_macaya_a549_endpoints import verify_reviewed_sources
    (tmp_path/'source_data.xlsx').write_bytes(b'changed source')
    with pytest.raises(ValueError, match='re-review'):
        verify_reviewed_sources(tmp_path)
