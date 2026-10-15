"""Recover source numeric flow endpoints without resolving conflicting gate definitions.

This is an external reference, not training labels for total Annexin positivity.
"""
import argparse
import hashlib
import json
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl

WORKBOOK_SHA = '41676635f1e897a8d262cc88b0cbc6878b390822a11830737cf7341a951f6b46'
DOSES = [0, 10, 20, 30, 50, 80]
# Personally read printed representative flow quadrants in Supplementary Fig S7C.
EARLY = {'A549': [2.96, 2.60, 4.20, 5.60, 14.5, 16.8],
         'NCI-H1299': [1.77, 2.47, 2.81, 5.05, 10.8, 21.6]}
DOUBLE = {'A549': [6.78, 5.66, 6.43, 7.36, 20.6, 26.6],
          'NCI-H1299': [2.76, 4.38, 6.07, 6.44, 15.9, 23.7]}


def extract(workbook):
    sheet = workbook['figS7']
    if sheet['C17'].value != 'figS7.b':
        raise ValueError('Reviewed panel binding changed')
    rows, representative = [], []
    for cell, row in [('A549', 19), ('NCI-H1299', 23)]:
        expected_name = 'A549' if cell == 'A549' else 'H1299'
        if sheet.cell(row-1, 3).value != expected_name:
            raise ValueError('Cell background binding changed')
        for i, dose in enumerate(DOSES):
            col = 4 + 3*i
            header = str(sheet.cell(row-1, col).value)
            if header.replace('μm', '').strip() != str(dose):
                raise ValueError('Dose binding changed')
            for j in range(3):
                source = sheet.cell(row, col+j)
                value = source.value
                if not isinstance(value, (int, float)) or not np.isfinite(value) or not 0 <= value <= 100:
                    raise ValueError('Invalid source percentage')
                rows.append(dict(condition_id=f'{cell}_capivasertib_{dose}uM_48h',cell_background=cell,
                    drug='capivasertib',dose_uM=dose,exposure_hours=48,measurement_hours=48,
                    source_panel='S7B',source_sheet='figS7',source_cell=source.coordinate,
                    source_numeric_index=j+1,value_percent=value,
                    endpoint_definition='reported_PI_Annexin_double_positive_axis',
                    methods_definition_conflict=True,total_annexin_training_eligible=False,
                    biological_repeat_id=None,technical_repeat_id=None,
                    observations_are_single_cells=False))
            if sheet.cell(row,col).value != DOUBLE[cell][i]:
                raise ValueError('Representative flow versus workbook first value changed')
            representative.append(dict(cell_background=cell,dose_uM=dose,source_panel='S7C',
                representative_double_positive_percent=DOUBLE[cell][i],
                representative_early_annexin_percent=EARLY[cell][i],
                representative_total_annexin_percent=DOUBLE[cell][i]+EARLY[cell][i],
                linked_to_independent_repeat=False))
    return pd.DataFrame(rows), pd.DataFrame(representative)


def recover(source, output):
    if output.exists():
        raise FileExistsError('Use a new versioned output directory')
    raw=source.read_bytes()
    if hashlib.sha256(raw).hexdigest()!=WORKBOOK_SHA:
        raise ValueError('Personally reviewed workbook changed')
    rows, representative=extract(openpyxl.load_workbook(source,data_only=True))
    output.mkdir(parents=True)
    rows.to_csv(output/'reported_flow_observations.csv',index=False)
    representative.to_csv(output/'representative_flow_quadrants.csv',index=False)
    summary=rows.groupby(['condition_id','cell_background','dose_uM']).value_percent.agg(['count','mean','std']).reset_index()
    summary.to_csv(output/'condition_summary.csv',index=False)
    reread=pd.read_csv(output/'reported_flow_observations.csv')
    if not np.array_equal(reread.value_percent.to_numpy(),rows.value_percent.to_numpy()):
        raise AssertionError('Numeric round trip failed')
    manifest=dict(doi='10.1002/advs.202523667',pmcid='PMC13335885',pmcaiid='13459879',
        source_workbook_sha256=WORKBOOK_SHA,conditions=len(summary),numeric_observations=len(rows),
        A549_conditions=6,A549_numeric_observations=18,exposure_hours=48,
        endpoint_conflict='Methods Annexin V positive; S7B axis and first source values correspond to double-positive quadrant. Do not relabel as total Annexin.',
        repeat_hierarchy='Three numbers per condition; Methods states three biological experiments with technical triplicates; workbook does not identify their aggregation or pairing.',
        sequence_status='Fig7K-O contains sequential protocol and TUNEL images, no corresponding numeric N/O table in the recovered workbook; no matched sequence RNA recovered.',
        RNA_status='Text describes FUCA2-knockdown RNA-seq; recovered workbook contains figure numbers, no expression matrix or accession identified.',
        total_annexin_training_eligible=False,complete_chain_validated=False,
        files={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in output.iterdir()})
    (output/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    return manifest

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();print(json.dumps(recover(a.workbook,a.output),indent=2))
