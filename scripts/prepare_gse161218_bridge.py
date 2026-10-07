"""Fetch limited official sources for the GSE161218 expression-apoptosis audit.

Dataset payloads remain under 08-datasets/GSE161218 and outside Git.
This does not download sequencing reads or reinterpret missing conditions.
"""

from __future__ import annotations

import hashlib
import json
import gzip
import io
from datetime import datetime, timezone
from pathlib import Path

import requests
import pandas as pd
import openpyxl
import fitz
from bs4 import BeautifulSoup


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "08-datasets" / "GSE161218"
RAW = DATA / "GSE161218" / "sources"
MAX_FILE_BYTES = 20_000_000
BASE = ("https://media.springernature.com/original/springer-static/esm/"
        "art%3A10.1038%2Fs41467-023-41828-z/MediaObjects/")


def fetch(url: str, filename: str, manifest: list[dict]) -> bytes:
    path = RAW / filename
    if path.exists():
        content = path.read_bytes()
        status = "reused_existing"
    else:
        response = requests.get(url, timeout=90, stream=True)
        response.raise_for_status()
        declared = response.headers.get("content-length")
        if declared and int(declared) > MAX_FILE_BYTES:
            raise ValueError(f"File exceeds bound: {filename}")
        chunks = []
        size = 0
        for chunk in response.iter_content(65_536):
            size += len(chunk)
            if size > MAX_FILE_BYTES:
                raise ValueError(f"File exceeds bound: {filename}")
            chunks.append(chunk)
        content = b"".join(chunks)
        path.write_bytes(content)
        status = "downloaded"
    manifest.append({
        "filename": filename, "url": url,
        "retrieved_or_verified_at_utc": datetime.now(timezone.utc).isoformat(),
        "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest(),
        "status": status,
    })
    return content


def extract_annexin() -> None:
    workbook = openpyxl.load_workbook(
        RAW / "41467_2023_41828_MOESM7_ESM.xlsx", read_only=True, data_only=True)
    sheet = workbook["Fig Suppl 2"]
    rows = []
    for row_start, lines in [(12, ["H1792", "H2009", "A549"]),
                             (19, ["H1568", "H1993", "H1437"])]:
        for line, column_start in zip(lines, [3, 9, 15]):
            for drug, offset, tram_dose, lest_dose in [
                ("Control", 0, 0, 0), ("Trametinib", 1, 0.5, 0),
                ("Lestaurtinib", 2, 0, 0.625), ("Simultaneous combination", 3, 0.5, 0.625),
            ]:
                for replicate in range(3):
                    cell = sheet.cell(row_start + replicate, column_start + offset)
                    assert isinstance(cell.value, (int, float)) and 0 <= cell.value <= 100
                    rows.append({
                        "cell_line": line, "condition": drug,
                        "trametinib_uM": tram_dose, "lestaurtinib_uM": lest_dose,
                        "time_h": 24, "experiment_number_within_endpoint": replicate + 1,
                        "annexin_v_positive_percent": cell.value,
                        "source_sheet": sheet.title, "source_cell": cell.coordinate,
                        "protocol_source": "Supplementary PDF page 5, Figure 2C legend",
                        "same_sample_as_rna": "not established",
                    })
    assert len(rows) == 72
    pd.DataFrame(rows).to_csv(DATA / "GSE161218" / "annexin_v_fig_s2c_long.csv", index=False)
    mido_rows = []
    sheet = workbook["Fig Suppl 3"]
    for row_start, lines in [(27, ["H1792", "H2009", "A549"]),
                             (34, ["H1568", "H1993", "H1437"])]:
        for line, column_start in zip(lines, [3, 9, 15]):
            for condition, offset, tram_dose, mido_dose in [
                ("Control", 0, 0, 0), ("Trametinib", 1, 0.5, 0),
                ("Midostaurin", 2, 0, 0.625), ("Simultaneous combination", 3, 0.5, 0.625),
            ]:
                for replicate in range(3):
                    cell = sheet.cell(row_start + replicate, column_start + offset)
                    assert cell.value is None or (isinstance(cell.value, (int, float)) and 0 <= cell.value <= 100)
                    mido_rows.append({
                        "source_block": "Fig_Suppl_3D_E", "cell_line": line,
                        "condition": condition, "trametinib_uM": tram_dose,
                        "midostaurin_uM": mido_dose, "time_h": 24,
                        "experiment_number_within_endpoint": replicate + 1,
                        "annexin_v_positive_percent": cell.value,
                        "value_status": "missing_in_source" if cell.value is None else "observed",
                        "source_sheet": sheet.title, "source_cell": cell.coordinate,
                        "protocol_source": "Supplementary PDF page 7, Figure 3D/E legend",
                        "same_sample_as_rna": "not established",
                    })
    assert len(mido_rows) == 72
    assert sum(row["value_status"] == "observed" for row in mido_rows) == 70
    pd.DataFrame(mido_rows).to_csv(DATA / "GSE161218" / "annexin_v_fig_s3de_midostaurin_long.csv", index=False)
    workbook.close()
    document = fitz.open(RAW / "41467_2023_41828_MOESM1_ESM.pdf")
    text = "\n\n".join(f"## PDF page {i+1}\n\n{page.get_text()}"
                          for i, page in enumerate(document))
    (DATA / "GSE161218" / "supplement_extracted.md").write_text(text, encoding="utf-8")


def verify_counts(rows: list[dict], manifest: list[dict]) -> None:
    metrics = []
    reference_genes = None
    for row in rows:
        url = row["supplementary_file_1"][0].replace("ftp://", "https://", 1)
        name = url.rsplit("/", 1)[1]
        content = fetch(url, name, manifest)
        table = pd.read_csv(io.BytesIO(gzip.decompress(content)), sep="\t", comment="#")
        assert "Geneid" in table.columns and len(table.columns) == 7, table.columns.tolist()
        genes = table["Geneid"].astype(str)
        counts = pd.to_numeric(table.iloc[:, -1], errors="raise")
        assert genes.is_unique and counts.notna().all() and (counts >= 0).all()
        assert (counts == counts.astype("int64")).all()
        if reference_genes is None:
            reference_genes = genes
        assert genes.equals(reference_genes), "Count gene order differs among samples"
        metrics.append({"GSM": row["GSM"], "title": row["title"][0],
                        "file": name, "gene_rows": len(table),
                        "positive_gene_rows": int((counts > 0).sum()),
                        "summed_gene_counts": int(counts.sum()),
                        "unique_genes": True, "integer_nonnegative": True})
    pd.DataFrame(metrics).to_csv(DATA / "GSE161218" / "counts_read_validation.csv", index=False)
    print(json.dumps(metrics, ensure_ascii=False, indent=2))


def write_audit_tables(rows: list[dict], manifest: list[dict]) -> None:
    endpoint = pd.read_csv(DATA / "GSE161218" / "annexin_v_fig_s2c_long.csv")
    sample_table = []
    for row in rows:
        title = row["title"][0]
        group, replicate = title.split("_")
        condition, tram, lest = {
            "CTL": ("DMSO control", 0, 0),
            "Lest": ("Lestaurtinib", 0, 0.625),
            "Tram": ("Trametinib", 0.5, 0),
        }[group]
        sample_table.append({
            "GSM": row["GSM"], "title": title, "cell_line": "H1792",
            "condition": condition, "trametinib_uM": tram, "lestaurtinib_uM": lest,
            "time_h": 16, "replicate_number_within_rna": int(replicate),
            "condition_basis": "GSM source_name/title + shared treatment protocol; paper Fig 1F control",
            "endpoint_replicate_link": "not established",
            "expression_file": Path(row["supplementary_file_1"][0]).name,
        })
    pd.DataFrame(sample_table).to_csv(DATA / "GSE161218" / "rna_sample_conditions.csv", index=False)
    lines = ["# GSE161218 表达—凋亡条件匹配核验", "",
             "> 核验日期：2026-09-18。源数据为公开实测结果；此处做数据审计，没有训练预测模型。", "",
             "## RNA 样本逐项登记", "",
             "所有 RNA 样本均为 H1792、16 h。组别由 GSM 的 title/source_name 与正文 Fig. 1F/G 共同确认。", "",
             "GSM 的 treatment_protocol 对全部九条记录复用同一段两种药物剂量说明，连 CTL 也如此；不能将它解析为九条组合处理。", "",
             "| GSM | 原始标题 | 条件 | Trametinib μM | Lestaurtinib μM | 组内重复编号 |", "|---|---|---|---:|---:|---:|"]
    for row in sample_table:
        lines.append(f"| {row['GSM']} | {row['title']} | {row['condition']} | {row['trametinib_uM']} | {row['lestaurtinib_uM']} | {row['replicate_number_within_rna']} |")
    lines += ["", "## Annexin V 阳性百分比原始值", "",
              "来源：官方 Source Data 工作表 `Fig Suppl 2`；处理均为 24 h。Tram=0.5 μM，Lest=0.625 μM，Combo 为这两剂量的同时双药处理。", "",
              "补充 PDF 第 5 页图注规定 n=3 independent experiments。每行的三个数按源表 Exp 1、Exp 2、Exp 3 排列；它们不自动对应 RNA 的 _1、_2、_3。", "",
              "| 细胞系 | 条件 | Exp 1 (%) | Exp 2 (%) | Exp 3 (%) | 源单元格 |", "|---|---|---:|---:|---:|---|"]
    for (line, condition), group in endpoint.groupby(["cell_line", "condition"], sort=False):
        values = group["annexin_v_positive_percent"].tolist()
        cells = ", ".join(group["source_cell"])
        lines.append(f"| {line} | {condition} | {values[0]:g} | {values[1]:g} | {values[2]:g} | {cells} |")
    lines += ["", "## 连接层级与允许的用途", "",
              "| 连接候选 | 同研究 | 同细胞系 | 同药物及剂量 | 同处理时长 | 同一样本/可配对重复 | 判断 |", "|---|---|---|---|---|---|---|",
              "| H1792 Tram RNA → Annexin | 是 | 是 | 是，0.5 μM | 否，16→24 h | 未建立 | 可作条件级、跨时间辅助验证；不是同终点配对标签 |",
              "| H1792 Lest RNA → Annexin | 是 | 是 | 是，0.625 μM | 否，16→24 h | 未建立 | 可作条件级、跨时间辅助验证；不是同终点配对标签 |",
              "| H1792 CTL RNA → Annexin Ctrl | 是 | 是 | RNA 图注为 DMSO；Annexin 图注为 untreated，溶剂一致性未明确 | 否，16→24 h | 未建立 | 单独标记对照定义缺口，不默认相同 |",
              "| H1792 Combo RNA → Annexin Combo | 是 | 是 | RNA 中没有组合组 | 不适用 | 不存在该 RNA 组 | 只有同时组合终点，不能构造对应 RNA |",
              "| A549 等其余五细胞系 RNA → Annexin | 是 | RNA 不含这些细胞系 | 未有对应 RNA | 不适用 | 不存在该 RNA 组 | 可保留外部终点；需其他表达来源另行匹配 |", "",
              "同研究并不建立个体级联合观测。只有 H1792 的三个名义条件在两种实验中都出现，其中两个单药条件的药物和剂量明确一致；各自组内三重复不能交叉扩增为独立配对训练样本。", "",
              "可支持：限定条件下的外部药效排序或生物学一致性检查，以及在显式建模 16 h 表达→24 h 结局假设下的探索性分析。不能单独支撑：通用高维终点评价器的可靠训练/独立校准、单细胞死亡状态监督、任意剂量外推或真实序贯验证。", "",
              "本次没有将 cleaved PARP Western blot、72 h viability 或肿瘤体积变化混入 Annexin V 百分比标签。", "",
              "## 原文定位", "",
              "- [GEO Series](https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE161218)：Overall design；各 GSM 的 source_name/title/treatment_protocol。",
              "- [论文](https://www.nature.com/articles/s41467-023-41828-z)：Results 中单药 RNAseq 段落；Fig. 1F/G；Methods 的 Apoptosis。",
              f"- [补充 PDF]({BASE}41467_2023_41828_MOESM1_ESM.pdf)：PDF 第 5 页，Supplementary Fig. 2C/D 图注。",
              f"- [源数据 XLSX]({BASE}41467_2023_41828_MOESM7_ESM.xlsx)：`Fig Suppl 2` C12:F14、I12:L14、O12:R14、C19:F21、I19:L21、O19:R21。", ""]
    (DATA / "CONDITION_MATCHING.md").write_text("\n".join(lines), encoding="utf-8")
    mido = pd.read_csv(DATA / "GSE161218" / "annexin_v_fig_s3de_midostaurin_long.csv")
    mido_lines = ["# Midostaurin 补充终点核验", "",
                  "> 独立保留 Supplementary Fig. 3D/E 的数据块，不能当作 GSE161218 中新增的 RNA 组。", "",
                  "来源：补充 PDF 第 7 页图注；官方 Source Data 的 `Fig Suppl 3`。所有处理为 24 h，Trametinib=0.5 μM，Midostaurin=0.625 μM。Combo 为同时组合，无给药先后顺序。", "",
                  "图注写 n=3 independent experiments，但源表 A549 的 Control Exp 3（O29）和 Tram Exp 3（P29）为空。因此共 72 个预期单元格、70 个实际数值，不能填零或伪造第三重复。", "",
                  "| 细胞系 | 条件 | Exp 1 (%) | Exp 2 (%) | Exp 3 (%) | 源单元格 |", "|---|---|---:|---:|---:|---|"]
    for (line, condition), group in mido.groupby(["cell_line", "condition"], sort=False):
        values = ["缺失" if pd.isna(v) else f"{v:g}" for v in group["annexin_v_positive_percent"]]
        mido_lines.append(f"| {line} | {condition} | {values[0]} | {values[1]} | {values[2]} | {', '.join(group['source_cell'])} |")
    controls_2 = endpoint[endpoint["condition"].eq("Control")]
    controls_3 = mido[mido["condition"].eq("Control")]
    joined = controls_2.merge(controls_3, on=["cell_line", "experiment_number_within_endpoint"], suffixes=("_s2", "_s3"))
    both = joined.dropna(subset=["annexin_v_positive_percent_s2", "annexin_v_positive_percent_s3"])
    identical = int((both["annexin_v_positive_percent_s2"] == both["annexin_v_positive_percent_s3"]).sum())
    mido_lines += ["", "## 对照复用检查与解释边界", "",
                   f"- 按细胞系和表内 Exp 编号比较两个附图的数据块，有 {len(both)} 个共同非空对照位置，其中 {identical} 个数值相同；未发现整组直接重复使用的对照数字。",
                   "- 数字不同不能证明跨附图实验在统计上完全独立；源表 Exp 1/2/3 不是全局样本编号。因此保留 source_block，不合并成六个独立重复，也不建立跨图个体配对。",
                   "- 两块 Trametinib 对照药的数值也并不相同，不能把来源块差异抹掉后宣称测得稳定的统一药效比例。",
                   "- GSE161218 无 Midostaurin RNA、无 Tram+Midostaurin RNA，且 RNA 仅在 H1792、16 h。新增的是外部终点参考，非新增配对表达监督。",
                   "- Annexin V 阳性百分比是论文在其门控与采样协议下的操作性终点，不自动等于初始培养细胞总体最终死亡比例。", "",
                   f"[补充 PDF 第 7 页]({BASE}41467_2023_41828_MOESM1_ESM.pdf)；[源数据 XLSX 的 Fig Suppl 3]({BASE}41467_2023_41828_MOESM7_ESM.xlsx)。", ""]
    (DATA / "MIDOSTAURIN_ENDPOINTS.md").write_text("\n".join(mido_lines), encoding="utf-8")
    manifest_lines = ["# 官方源文件清单与校验值", "",
                      "> 获取/复核时间以 UTC 记录。SOFT/HTML URL 没有固定下载文件名，因此使用 accession/DOI 命名；官方附件与 counts 保留原始文件名。", "",
                      "| 文件与来源 | 字节数 | SHA-256 | 获取或复核时间 UTC |", "|---|---:|---|---|"]
    for item in manifest:
        manifest_lines.append(f"| [{item['filename']}]({item['url']}) | {item['bytes']} | `{item['sha256']}` | {item['retrieved_or_verified_at_utc']} |")
    (DATA / "SOURCE_MANIFEST.md").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")


def soft_values(text: str, prefix: str) -> list[str]:
    return [line.split(" = ", 1)[1] for line in text.splitlines()
            if line.startswith(prefix + " = ")]


def main() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    manifest: list[dict] = []
    query = "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc="
    suffix = "&targ=self&form=text&view=full"
    series = fetch(query + "GSE161218" + suffix, "GSE161218.soft", manifest)
    samples = soft_values(series.decode("utf-8"), "!Series_sample_id")
    rows = []
    for sample in samples:
        text = fetch(query + sample + suffix, sample + ".soft", manifest).decode("utf-8")
        row = {"GSM": sample}
        for field in ["title", "source_name_ch1", "characteristics_ch1",
                      "treatment_protocol_ch1", "growth_protocol_ch1",
                      "data_processing", "supplementary_file_1", "relation"]:
            row[field] = soft_values(text, "!Sample_" + field)
        rows.append(row)
    (DATA / "GSE161218" / "sample_records.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    html = fetch("https://www.nature.com/articles/s41467-023-41828-z",
                 "s41467-023-41828-z.html", manifest)
    soup = BeautifulSoup(html, "html.parser")
    paragraphs = [element.get_text(" ", strip=True)
                  for element in soup.find_all(["h1", "h2", "h3", "p", "figcaption"])]
    (DATA / "GSE161218" / "article_extracted.md").write_text(
        "# Publisher HTML text extraction\n\n" + "\n\n".join(paragraphs), encoding="utf-8")
    for filename in ["41467_2023_41828_MOESM1_ESM.pdf",
                     "41467_2023_41828_MOESM7_ESM.xlsx"]:
        fetch(BASE + filename, filename, manifest)
    extract_annexin()
    verify_counts(rows, manifest)
    (DATA / "GSE161218" / "source_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    write_audit_tables(rows, manifest)
    print(json.dumps({"samples": len(rows), "files": len(manifest),
                      "bytes": sum(item["bytes"] for item in manifest)}, indent=2))


if __name__ == "__main__":
    main()
