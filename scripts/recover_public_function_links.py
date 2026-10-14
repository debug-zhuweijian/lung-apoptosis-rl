"""Recover original functional observations and condition-level RNA links.

No biological replicates are invented, no source is converted to a different
endpoint definition, and no learning roles or clinical efficacy are certified.
"""
from __future__ import annotations
import argparse
import gzip
from io import StringIO
import json
from pathlib import Path
import re
import sys
import numpy as np
import pandas as pd
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT)); sys.path.insert(0,str(ROOT/'src'))
from lung.data.public_source_recovery import (file_sha256, load_expression_packet,
    read_geo_samples, validate_function_observations, verify_source_lock, write_expression_packet)
from scripts import prepare_gse161218_bridge as macaya
from scripts import align_sarin_state_expression as sarin


def save_endpoints(folder: Path, observations: pd.DataFrame, links: pd.DataFrame) -> dict:
    validate_function_observations(observations)
    observations.to_csv(folder/'functional_observations.csv',index=False)
    links.to_csv(folder/'condition_level_links.csv',index=False)
    reread=pd.read_csv(folder/'functional_observations.csv')
    validate_function_observations(reread)
    np.testing.assert_allclose(reread.value,observations.value,rtol=1e-14,atol=1e-14,equal_nan=True)
    if not links.link_id.is_unique: raise ValueError('Condition links must be unique')
    known=set(observations.observation_id)
    for row in links.to_dict('records'):
        ids=json.loads(row['endpoint_observation_ids_json'])
        if len(ids)!=len(set(ids)) or not set(ids).issubset(known):
            raise ValueError('Condition link contains duplicated or unknown endpoint identities')
        if row['pairing_level'] not in ['cross_study_condition_correspondence',
                'same_study_condition_early_RNA_later_endpoint', 'external_endpoint_without_RNA']:
            raise ValueError('Unsupported pairing claim')
    result={'source_positions':len(observations),'observed_values':int(observations.value.notna().sum()),
        'missing_positions':int(observations.value.isna().sum()),
        'endpoint_domains':sorted(observations.endpoint_domain.unique()),
        'condition_link_count':len(links),'whole_project_qualified':False,
        'learning_roles_assigned':False,'model_effect_validated':False,
        'files':{name:file_sha256(folder/name) for name in ['functional_observations.csv','condition_level_links.csv']}}
    (folder/'functional_receipt.json').write_text(json.dumps(result,indent=2))
    return result


def recover_macaya(source: Path, expression: Path, output: Path) -> dict:
    folder=output/'Macaya2023'; folder.mkdir()
    # Reuse the original source-cell extractor; do not run its historical main,
    # which also rewrites documentation and downloads into an old local layout.
    macaya.RAW=source/'macaya'; macaya.DATA=folder
    (folder/'GSE161218').mkdir()
    macaya.extract_annexin()
    rna=load_expression_packet(expression/'GSE161218',expected_scale='bulk_raw_counts',expected_background='H1792')['samples']
    rows=[]
    for block,filename,second in [
        ('Fig_Suppl_2C','annexin_v_fig_s2c_long.csv','Lestaurtinib'),
        ('Fig_Suppl_3DE','annexin_v_fig_s3de_midostaurin_long.csv','Midostaurin')]:
        raw=pd.read_csv(folder/'GSE161218'/filename)
        for row in raw.to_dict('records'):
            agents=[]
            if row['trametinib_uM']>0: agents.append({'name':'Trametinib','dose':row['trametinib_uM'],'unit':'uM'})
            key=second.lower()+'_uM'
            if row[key]>0: agents.append({'name':second,'dose':row[key],'unit':'uM'})
            pos=row['source_sheet']+'!'+row['source_cell']
            rows.append(dict(observation_id='Macaya2023:'+pos,source_position=pos,source_id='Macaya2023',
                source_file='41467_2023_41828_MOESM7_ESM.xlsx',source_sha256=file_sha256(macaya.RAW/'41467_2023_41828_MOESM7_ESM.xlsx'),
                cell_line=row['cell_line'],background_id='Macaya2023:'+row['cell_line'],
                condition=row['condition'],experiment_block=block,
                replicate_index=row['experiment_number_within_endpoint'],
                replicate_type='independent experiment, supplement figure legend; cross-panel identity not asserted',
                agents_json=json.dumps(agents,sort_keys=True),treatment_order='simultaneous' if len(agents)>1 else 'single_or_control',
                exposure_duration_h=24,measurement_time_h=24,
                control_definition='untreated per figure legend; not certified equal to RNA DMSO',
                endpoint_domain='annexin_positive_fraction',raw_unit='percent',raw_value=row['annexin_v_positive_percent'],
                value=row['annexin_v_positive_percent']/100,
                value_status='observed' if pd.notna(row['annexin_v_positive_percent']) else 'missing_in_source',
                denominator='FSC/SSC alive and single-cell gated acquired population; see supplement Fig2D/3E',
                source_protocol_locator=row['protocol_source'],role='unassigned',paired_to_rna_replicate=False,
                plate_id=None,batch_id=None,biological_repeat_cross_assay_id=None))
    obs=pd.DataFrame(rows);links=[]
    for (block,line,condition),group in obs.groupby(['experiment_block','cell_line','condition'],sort=False):
        # No Cartesian product between the independent RNA and functional replicates.
        suffix={'Control':'CTL','Trametinib':'Tram','Lestaurtinib':'Lest'}.get(condition)
        eligible=line=='H1792' and block=='Fig_Suppl_2C' and suffix is not None
        rnaids=rna.loc[rna.source_title.str.startswith(suffix+'_'),'sample_id'].tolist() if eligible else []
        refs=rna.loc[rna.source_title.str.startswith('CTL_'),'sample_id'].tolist() if eligible else []
        links.append(dict(link_id=f'Macaya2023:{block}:{line}:{condition}',cell_line=line,
            endpoint_observation_ids_json=json.dumps(group.observation_id.tolist()),
            rna_sample_ids_json=json.dumps(rnaids),reference_rna_sample_ids_json=json.dumps(refs),
            pairing_level='same_study_condition_early_RNA_later_endpoint' if eligible else 'external_endpoint_without_RNA',
            rna_time_h=16 if eligible else None,function_time_h=24,role='unassigned',
            missing_connection='different experiments and time; control equivalence not established' if eligible else 'no RNA in this source for this condition',
            exact_protocol_equivalence=False,paired_rna_function_replicates=False))
    result=save_endpoints(folder,obs,pd.DataFrame(links))
    result['A549_observed_values']=int((obs.cell_line.eq('A549') & obs.value.notna()).sum())
    result['H1792_RNA_condition_links']=sum(bool(json.loads(r['rna_sample_ids_json'])) for r in links)
    return result


def recover_sarin(source: Path, followup: Path, expression: Path, output: Path) -> dict:
    folder=output/'Sarin2017_2018'; folder.mkdir()
    rawpath=source/'GSE108214/GSE108214_family.soft.gz'
    text=gzip.open(rawpath,'rt',encoding='utf-8').read()
    platform=pd.read_csv(StringIO(text.split('!platform_table_begin\n',1)[1].split('!platform_table_end',1)[0]),sep='\t',dtype=str,keep_default_na=False)
    records=read_geo_samples(rawpath); matrices=[]; names=[]; samples=[]; probes=None
    for block in re.split(r'(?=\^SAMPLE = )',text)[1:]:
        name=block.splitlines()[0].split(' = ',1)[1]
        table=pd.read_csv(StringIO(block.split('!sample_table_begin\n',1)[1].split('!sample_table_end',1)[0]),sep='\t',dtype=str,keep_default_na=False)
        if probes is None: probes=table.ID_REF.tolist()
        if table.ID_REF.tolist()!=probes or not table.ID_REF.is_unique: raise ValueError('Sarin probe/sample identity mismatch')
        matrices.append(pd.to_numeric(table.VALUE,errors='raise').to_numpy(dtype=np.float64));names.append(name)
        record=records[name]; chars=dict(v.split(': ',1) for v in record['characteristics_ch1'])
        if chars['sensitivity'] not in ['sensitive','resistant']: raise ValueError('Unknown Sarin background')
        state='A549' if chars['sensitivity']=='sensitive' else 'A549rCDDP2000'
        dose=float(chars['treatment conc'])
        samples.append(dict(sample_id=name,source_column=name,source_title=record['title'][0],cell_line='A549',
            background_id=state,dose_uM=dose,measurement_time_h=24,drug='Cisplatin' if dose else 'untreated',
            role='unassigned',paired_to_function_replicate=False,observed_pretreatment_cell=False,
            maintenance_cisplatin_ug_ml=2 if state!='A549' else None,
            withdrawal_before_challenge_h=None,replicate_scope='RNA suffix; not paired to functional replicates',
            source_fields_json=json.dumps(record,sort_keys=True)))
    values=np.stack(matrices,axis=1); sampleframe=pd.DataFrame(samples)
    annotation=platform.set_index('ID').loc[probes].reset_index().rename(columns={'ID':'probe_id'})
    axis=pd.read_parquet(expression/'gene_axis.parquet')
    wanted=set().union(*(set(sarin.transcript_ids(r)) for r in annotation.to_dict('records')))
    tx=sarin.read_transcripts(followup/'annotation/Homo_sapiens.GRCh37.75.gtf.gz',wanted)
    hg=pd.read_csv(followup/'annotation/hgnc_complete_set.txt',sep='\t',dtype=str,keep_default_na=False)
    ledger=sarin.mapping_ledger(annotation,axis,hg,tx)
    aligned,mask=sarin.aligned_matrix(values,ledger,len(axis))
    chosen=ledger.gene_matrix_status.eq('unique_resolved_probe')
    frame=pd.DataFrame(values[chosen],columns=names)
    frame.insert(0,'ensembl_id',ledger.loc[chosen,'resolved_ensembl_gene_id'].to_numpy())
    packet=write_expression_packet(folder/'RNA',frame,axis,sampleframe,identifier='ensembl_id',scale='bulk_microarray_processed',
        metadata={'source_id':'GSE108214','native_domain':'background-corrected, quantile-normalized deposited microarray VALUE',
        'log_base':None,'probe_selection':'only one verified probe per gene; no averaging, full probe values kept separately',
        'platform':'GPL17077, GRCh37/hg19; original annotation release unspecified',
        'transcript_bridge':'Ensembl release75 build-compatible, NOT original annotation certified',
        'identity_bridge':'current downloaded HGNC snapshot; NOT the old frozen crosswalk',
        'functional_pairing':'independent studies, condition-level correspondence only'})
    reread=load_expression_packet(folder/'RNA',expected_scale='bulk_microarray_processed',expected_gene_ids=axis.ensembl_id.tolist())
    np.testing.assert_array_equal(reread['expression'],aligned)
    np.savez_compressed(folder/'all_original_probes.npz',expression=values,probe_ids=np.array(probes,dtype=str),sample_ids=np.array(names,dtype=str))
    with np.load(folder/'all_original_probes.npz',allow_pickle=False) as arr:
        np.testing.assert_array_equal(arr['expression'],values)
    ledger.to_csv(folder/'full_probe_mapping.csv.gz',index=False)
    tx.to_csv(folder/'transcript_annotation_subset.csv.gz',index=False)
    docpath=followup/'sarin/pone.0181081.s008.docx'; doc=Document(docpath)
    table=doc.tables[5]
    if 'FITC-Annexin' not in table.cell(2,0).text or 'fold change' not in table.cell(2,0).text:
        raise ValueError('Sarin source table no longer has the expected endpoint heading')
    if [table.cell(0,j).text for j in [1,2,3]]!=['A549','A549rCDDP2000','A549rCDDP2000']:
        raise ValueError('Sarin source background header changed')
    rows=[];links=[]
    for col,state,dose in [(1,'A549',11),(2,'A549rCDDP2000',11),(3,'A549rCDDP2000',34)]:
        if table.cell(1,col).text!=f'{dose} µM': raise ValueError('Sarin source dose header changed')
        ids=[]
        for row in range(3,7):
            text=table.cell(row,col).text.strip()
            if not text: continue
            pos=f'S8:table6:R{row+1}:C{col+1}'; obsid='Sarin2017:'+pos;ids.append(obsid)
            rows.append(dict(observation_id=obsid,source_position=pos,source_id='Sarin2017',source_file=docpath.name,
                source_sha256=file_sha256(docpath),cell_line='A549',background_id=state,drug='Cisplatin',dose_uM=dose,
                exposure_duration_h=24,measurement_time_h=24,replicate_index=row-2,
                replicate_scope='published individual endpoint value; no shared RNA replicate ID',
                endpoint_domain='annexin_relative_fold',raw_unit='fold',raw_value=float(text),value=float(text),value_status='observed',
                denominator='respective untreated background Annexin reference; absolute fraction not supplied',
                role='unassigned',paired_to_rna_replicate=False,plate_id=None,batch_id=None,
                source_protocol_locator='Sarin2017 sec004/sec009/Fig5a; S8 table6; Sarin2018 methods4.2-4.3'))
        rna=sampleframe.loc[sampleframe.background_id.eq(state)&sampleframe.dose_uM.eq(dose),'sample_id'].tolist()
        refs=sampleframe.loc[sampleframe.background_id.eq(state)&sampleframe.dose_uM.eq(0),'sample_id'].tolist()
        links.append(dict(link_id=f'Sarin:{state}:{dose}uM:24h',cell_line='A549',background_id=state,
            endpoint_observation_ids_json=json.dumps(ids),rna_sample_ids_json=json.dumps(rna),
            reference_rna_sample_ids_json=json.dumps(refs),pairing_level='cross_study_condition_correspondence',
            rna_time_h=24,function_time_h=24,role='unassigned',exact_protocol_equivalence=False,
            paired_rna_function_replicates=False,
            missing_connection='different studies/cultures; batch and additional maintenance washout unknown'))
    result=save_endpoints(folder,pd.DataFrame(rows),pd.DataFrame(links))
    result.update(RNA_packet=packet,original_probe_values=int(values.size),source_probes=len(probes),
        mapping_status_counts=ledger.gene_matrix_status.value_counts().to_dict(),
        annotation_hashes={p.name:file_sha256(p) for p in [rawpath,followup/'annotation/hgnc_complete_set.txt',followup/'annotation/Homo_sapiens.GRCh37.75.gtf.gz',docpath]},
        full_probe_hash=file_sha256(folder/'all_original_probes.npz'),mapping_hash=file_sha256(folder/'full_probe_mapping.csv.gz'))
    return result


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root',type=Path,required=True)
    parser.add_argument('--followup-root',type=Path,required=True)
    parser.add_argument('--expression-root',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    source=args.source_root.resolve(); follow=args.followup_root.resolve(); out=args.output.resolve()
    lock=json.loads((ROOT/'configs/public_alignment_sources_v1.json').read_text())
    verify_source_lock(source,lock['sources'])
    lock2=json.loads((ROOT/'configs/public_alignment_followup_v1.json').read_text())
    verify_source_lock(follow,lock2['sources'])
    if out.exists(): raise FileExistsError('Use a new functional recovery namespace')
    out.mkdir(parents=True)
    results={'Macaya2023':recover_macaya(source,args.expression_root,out),
             'Sarin2017_2018':recover_sarin(source,follow,args.expression_root,out)}
    results['whole_project_qualified']=False
    results['function_models_trained']=False
    results['implementation_sha256']=file_sha256(Path(__file__))
    results['sources_locked']=True
    (out/'recovery_receipt.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
    print(json.dumps(results,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
