"""Restore a known core A549 functional source into the existing endpoint reader.

This reconstitutes existing observations; it must not be concatenated as a new
independent dataset. GEO RNA is H1792 at 16h, not A549 at the 24h endpoint.
"""
from __future__ import annotations
import argparse
import json
import shutil
from pathlib import Path
import sys
import numpy as np
import openpyxl
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'src'))
from scripts.build_endpoint_alignment_v1 import EndpointBuilder,write_alignment
from scripts.recover_public_expression import digest,read_soft
from lung.data.endpoint_alignment import load_alignment


REVIEWED_SOURCE_SHA256 = {
    "source_data.xlsx": "3160a2b9c91d2b3eafaba8de055a752556145445220d1cbc135b8609b8d630b4",
    "supplement.pdf": "902297818c2e6335961b762807141cba6171a9354df7980180f8c7fe9427c3c9",
}


def verify_reviewed_sources(directory):
    # Dose/time/gating were manually verified against this exact supplement.
    # Changed publications must be re-reviewed, not silently given old units.
    for name, expected in REVIEWED_SOURCE_SHA256.items():
        if digest(directory / name) != expected:
            raise ValueError(f"Source revision requires protocol re-review: {name}")


def extract_a549(book):
    rows=[]
    for sheet,start,block,second,header in [('Fig Suppl 2',12,'S2C','Lestaurtinib','Lest'),('Fig Suppl 3',27,'S3DE','Midostaurin','Mido')]:
        s=book[sheet]
        if s.cell(start-2,15).value!='A549 - Annexin V positive cells (%)':raise ValueError('A549 source heading mismatch')
        expected=['Control' if block=='S2C' else 'Ctrl','Tram',header,'Combo']
        if [s.cell(start-1,15+i).value for i in range(4)]!=expected:raise ValueError('Treatment column order changed')
        for offset,label in enumerate(expected):
            components=[]
            if offset in [1,3]:components.append(('Trametinib',.5))
            if offset in [2,3]:components.append((second,.625))
            for rep in range(3):
                cell=s.cell(start+rep,15+offset);v=cell.value
                if v is not None and (not isinstance(v,(int,float)) or not np.isfinite(v) or not 0<=v<=100):
                    raise ValueError('Invalid source percentage')
                rows.append(dict(block=block,source_sheet=sheet,source_cell=cell.coordinate,replicate=rep+1,
                                 components=components,percent=v,missing=v is None))
    return rows


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    for n in ['source','output']:ap.add_argument('--'+n,type=Path,required=True)
    a=ap.parse_args()
    a.source=a.source.resolve();a.output=a.output.resolve()
    if a.output.exists():raise FileExistsError('New output directory required')
    if not a.output.resolve().is_relative_to(ROOT):raise ValueError('Existing endpoint reader requires a project-relative source directory')
    verify_reviewed_sources(a.source)
    _,rna=read_soft(a.source/'geo.soft.gz')
    if len(rna)!=9 or any('H1792' not in ' '.join(m['source_name_ch1']) or '16 hours' not in ' '.join(m['treatment_protocol_ch1']) for _,m,_ in rna):
        raise ValueError('RNA background or clock changed: re-review required')
    book=openpyxl.load_workbook(a.source/'source_data.xlsx',read_only=True,data_only=True);rows=extract_a549(book);book.close()
    a.output.mkdir(parents=True);sources=a.output/'sources';sources.mkdir()
    for n in ['source_data.xlsx','supplement.pdf','paper.xml','geo.soft.gz']:
        shutil.copyfile(a.source/n,sources/n)
        if digest(a.source/n)!=digest(sources/n):raise ValueError('Source copy changed')
    b=EndpointBuilder();source=sources/'source_data.xlsx'
    for block in ['S2C','S3DE']:
        b.protocols['Macaya2023:'+block]=dict(source='https://doi.org/10.1038/s41467-023-41828-z',
            endpoint='Annexin_positive_fraction',denominator='events retained by source flow gate, not initial culture count',
            sampling_gate='FSC/SSC alive and singlet selection; Annexin V/7AAD',
            control_definition='untreated; not asserted equivalent to RNA DMSO control',
            replicate_type='source reports three independent experiments; missing source cells retained',
            source_RNA_background='H1792',source_RNA_exposure_h=16,A549_RNA_available=False,
            source_supplement_sha256=digest(sources/'supplement.pdf'))
    for r in rows:
        cid=b.condition('Macaya2023','Macaya2023:'+r['block'],r['block'],'A549',
                        'A549_parental_unmeasured_reference',r['components'],treatment_hours=24.)
        b.observation(cid,'Annexin_positive_fraction',None if r['missing'] else r['percent']/100,
                      r['replicate'],source,r['source_sheet']+'!'+r['source_cell'],
                      'source_cell_missing' if r['missing'] else '', 'missing' if r['missing'] else 'observed')
    write_alignment(b,a.output)
    loaded=load_alignment(a.output,ROOT)
    observations=loaded['observations']
    if len(loaded['conditions']) != 8 or len(observations) != 24:
        raise ValueError('Existing-reader condition or observation count mismatch')
    if observations.value.notna().sum() != 22 or observations.value.isna().sum() != 2:
        raise ValueError('Existing-reader missingness mismatch')
    for r in rows:
        found=observations[observations.source_locator.eq(r['source_sheet']+'!'+r['source_cell'])]
        if len(found)!=1 or (not r['missing'] and float(found.iloc[0].value)!=r['percent']/100):
            raise ValueError('Source-to-reader numeric identity mismatch')
    receipt=dict(known_core_source_reconstituted=True,new_independent_dataset=False,
                 observation_slots=len(rows),numeric_observations=sum(not r['missing'] for r in rows),
                 missing_observations=sum(r['missing'] for r in rows),source_condition_blocks=8,
                 distinct_intervention_sets=6,existing_endpoint_reader_readback=True,
                 formal_training=False,full_research_qualified=False,script_sha256=digest(__file__),
                 source_hashes={f.name:digest(f) for f in sources.iterdir()},
                 alignment_manifest_sha256=digest(a.output/'manifest.json'))
    (a.output/'recovery_receipt.json').write_text(json.dumps(receipt,indent=2));print(json.dumps(receipt,indent=2))


if __name__=='__main__':main()
